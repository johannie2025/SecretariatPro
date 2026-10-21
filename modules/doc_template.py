"""Modèles intelligents : détecte [Variable] et « Intitulé : ______ » puis remplit le document."""
import re
from pathlib import Path

BRACKET = re.compile(r"\[([^\[\]\n]{1,40})\]")
BLANK = re.compile(r"(?P<label>[^\n:_\.]{1,60}?)\s*:\s*(?P<blank>[_\.…]{4,})")


def _iter_paragraphs(container):
    for p in container.paragraphs:
        yield p
    for t in getattr(container, "tables", []):
        for row in t.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _all_paragraphs(doc):
    yield from _iter_paragraphs(doc)
    for s in doc.sections:
        for part in (s.header, s.footer):
            yield from _iter_paragraphs(part)


def detect_fields(path: str):
    """Retourne la liste ordonnée des champs : [{'id','label','kind'}]."""
    path = Path(path)
    if path.suffix.lower() == ".docx":
        from docx import Document
        texts = [p.text for p in _all_paragraphs(Document(str(path)))]
    elif path.suffix.lower() == ".txt":
        texts = path.read_text(encoding="utf-8", errors="replace").splitlines()
    else:
        raise ValueError("Format non pris en charge (utilisez .docx ou .txt).")
    seen, fields = set(), []
    for line in texts:
        for m in BRACKET.finditer(line):
            name = m.group(1).strip()
            if name not in seen:
                seen.add(name)
                fields.append({"id": name, "label": name, "kind": "bracket"})
        for m in BLANK.finditer(line):
            label = m.group("label").strip()
            if label and label not in seen:
                seen.add(label)
                fields.append({"id": label, "label": label, "kind": "blank"})
    return fields


def _transform(text: str, values: dict) -> str:
    def br(m):
        v = values.get(m.group(1).strip())
        return v if v else m.group(0)

    def bl(m):
        v = values.get(m.group("label").strip())
        return f"{m.group('label')} : {v}" if v else m.group(0)

    return BLANK.sub(bl, BRACKET.sub(br, text))


def _fix_paragraph(p, values):
    full = p.text
    new = _transform(full, values)
    if new == full:
        return
    runs = p.runs
    per_run = "".join(_transform(r.text, values) for r in runs)
    if per_run == new:           # remplacement sans perdre la mise en forme
        for r in runs:
            r.text = _transform(r.text, values)
    elif runs:                   # balise coupée sur plusieurs « runs » : on fusionne
        runs[0].text = new
        for r in runs[1:]:
            r.text = ""


def fill_template(src: str, dest: str, values: dict):
    src, dest = Path(src), Path(dest)
    values = {k: str(v) for k, v in values.items() if str(v).strip()}
    if src.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(str(src))
        for p in _all_paragraphs(doc):
            _fix_paragraph(p, values)
        doc.save(str(dest))
    else:
        text = src.read_text(encoding="utf-8", errors="replace")
        dest.write_text("\n".join(_transform(l, values) for l in text.splitlines()), encoding="utf-8")
    return dest
