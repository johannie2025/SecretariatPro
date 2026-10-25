"""OCR avec conservation de la mise en page : textes, tableaux, images/dessins.

Sorties : Word reconstruit (.docx), PDF « recherchable » fidèle à l'original, texte brut.
100 % local (RapidOCR + OpenCV), aucune donnée envoyée sur Internet.
"""
import io
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

from ._lazy import fitz

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
MAX_BYTES = 60 * 1024 * 1024
MAX_PAGES = 100
MAX_SIDE = 3600
PDF_ZOOM = 2.0
_engine = None


@dataclass
class Page:
    jpg: bytes
    w: int
    h: int
    dpi: float
    boxes: list = field(default_factory=list)
    blocks: list = field(default_factory=list)


@dataclass
class OcrResult:
    path: str
    kind: str                      # image | pdf_text | pdf_scan
    pages: list = field(default_factory=list)
    text: str = ""
    n_tables: int = 0
    n_figures: int = 0
    n_pages: int = 0
    figures: list = field(default_factory=list)      # [{'name','data'(png ou jpg),'page'}]
    ai_text: str = ""                                 # texte lu par l'OCR IA (si activé)


def _engine_get():
    global _engine
    if _engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _engine = RapidOCR()
    return _engine


# ------------------------------------------------------------------ détection
def ocr_boxes(arr):
    result, _ = _engine_get()(arr)
    out = []
    for box, text, score in result or []:
        xs = [p[0] for p in box]
        ys = [p[1] for p in box]
        if str(text).strip():
            out.append({"x0": min(xs), "y0": min(ys), "x1": max(xs), "y1": max(ys), "text": str(text).strip()})
    return out


def _cluster(values, gap=6):
    out, cur = [], []
    for v in sorted(values):
        if cur and v - cur[-1] > gap:
            out.append(sum(cur) / len(cur))
            cur = []
        cur.append(v)
    if cur:
        out.append(sum(cur) / len(cur))
    return out


def detect_tables(arr):
    import cv2
    gray = cv2.cvtColor(arr, cv2.COLOR_BGR2GRAY)
    h, w = gray.shape
    bw = cv2.adaptiveThreshold(cv2.bitwise_not(gray), 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY, 15, -2)
    kw, kh = max(25, w // 40), max(25, h // 40)
    hl = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (kw, 1)))
    vl = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, kh)))
    mask = cv2.dilate(cv2.bitwise_or(hl, vl), np.ones((3, 3), np.uint8), iterations=2)
    cnts, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    tables = []
    for c in cnts:
        x, y, cw, ch = cv2.boundingRect(c)
        if cw < 0.25 * w or ch < 0.03 * h:
            continue
        rows = _cluster([y + i for i, s in enumerate(hl[y:y + ch, x:x + cw].sum(axis=1)) if s > 0.5 * cw * 255])
        cols = _cluster([x + i for i, s in enumerate(vl[y:y + ch, x:x + cw].sum(axis=0)) if s > 0.5 * ch * 255])
        if len(rows) >= 2 and len(cols) >= 2:
            tables.append({"bbox": (x, y, x + cw, y + ch), "ys": rows, "xs": cols})
    return tables


def _page_background(gray):
    h, w = gray.shape
    m = max(6, min(h, w) // 60)
    edge = np.concatenate([gray[:m].ravel(), gray[-m:].ravel(), gray[:, :m].ravel(), gray[:, -m:].ravel()])
    return float(np.median(edge))


def _merge_rects(rects):
    rects = [list(r) for r in rects]
    changed = True
    while changed:
        changed = False
        for i in range(len(rects)):
            for j in range(i + 1, len(rects)):
                a, b = rects[i], rects[j]
                if a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]:      # chevauchement
                    rects[i] = [min(a[0], b[0]), min(a[1], b[1]), max(a[2], b[2]), max(a[3], b[3])]
                    del rects[j]
                    changed = True
                    break
            if changed:
                break
    return [tuple(r) for r in rects]


def _absorb_logo_text(full, boxes, rects, w, h):
    """Un logo = dessin + quelques lignes de texte : on garde l'ensemble comme une seule image."""
    import cv2
    if not rects:
        return rects
    k = max(25, w // 28)
    closed = cv2.morphologyEx(full, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (k, k)))
    _, labels, stats, _ = cv2.connectedComponentsWithStats(closed)
    out = []
    for (x0, y0, x1, y1) in rects:
        region = labels[y0:y1, x0:x1].ravel()
        region = region[region > 0]
        if not region.size:
            out.append((x0, y0, x1, y1))
            continue
        lab = int(np.bincount(region).argmax())
        sx, sy, sw, sh = (int(v) for v in stats[lab][:4])
        U = (sx, sy, sx + sw, sy + sh)
        inside = [b for b in boxes if _center_in(b, U)]
        share = (x1 - x0) * (y1 - y0) / float(sw * sh)          # part du dessin dans l'ensemble
        if (len(inside) <= 6 and sw * sh <= 0.25 * w * h and share >= 0.12 and
                U[0] <= x0 + 20 and U[1] <= y0 + 20 and U[2] >= x1 - 20 and U[3] >= y1 - 20):
            out.append((max(sx - 6, 0), max(sy - 6, 0), min(sx + sw + 6, w), min(sy + sh + 6, h)))
        else:
            out.append((x0, y0, x1, y1))
    return out


def detect_figures(arr, boxes, tables):
    """Photos, logos, dessins, cachets, signatures : tout ce qui n'est ni du texte lu ni un tableau."""
    import cv2
    h, w = arr.shape[:2]
    blur = cv2.GaussianBlur(arr, (3, 3), 0)
    gray = cv2.cvtColor(blur, cv2.COLOR_BGR2GRAY)
    bg = _page_background(gray)
    mask = (np.abs(gray.astype(np.int16) - bg) > 30)
    hsv = cv2.cvtColor(blur, cv2.COLOR_BGR2HSV)
    mask |= (hsv[:, :, 1] > 55) & (hsv[:, :, 2] > 50)          # zones colorées (logos, tampons)
    mask = mask.astype(np.uint8) * 255
    full = mask.copy()                                          # contenu complet (texte inclus), pour les logos
    for b in boxes:
        cv2.rectangle(mask, (int(b["x0"]) - 4, int(b["y0"]) - 4), (int(b["x1"]) + 4, int(b["y1"]) + 4), 0, -1)
    for t in tables:
        x0, y0, x1, y1 = t["bbox"]
        cv2.rectangle(mask, (x0 - 4, y0 - 4), (x1 + 4, y1 + 4), 0, -1)
    k = max(13, w // 55)
    closed = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (k, k)))
    closed = cv2.dilate(closed, np.ones((5, 5), np.uint8), iterations=2)
    cnts, _ = cv2.findContours(closed, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    rects = []
    for c in cnts:
        x, y, cw, ch = cv2.boundingRect(c)
        if cw < 0.05 * w or ch < 0.03 * h or cw * ch > 0.75 * w * h:
            continue
        if max(cw / ch, ch / cw) > 12:                       # simple trait / filet
            continue
        ink = cv2.countNonZero(mask[y:y + ch, x:x + cw])
        if ink < 150 or ink / float(cw * ch) < 0.025:            # trop clairsemé : bruit
            continue
        rects.append((max(x - 6, 0), max(y - 6, 0), min(x + cw + 6, w), min(y + ch + 6, h)))
    return _merge_rects(_absorb_logo_text(full, boxes, _merge_rects(rects), w, h))


# ------------------------------------------------------------------ structure
def _center_in(b, rect):
    cx, cy = (b["x0"] + b["x1"]) / 2, (b["y0"] + b["y1"]) / 2
    return rect[0] <= cx <= rect[2] and rect[1] <= cy <= rect[3]


def group_lines(boxes):
    items = sorted(boxes, key=lambda b: ((b["y0"] + b["y1"]) / 2, b["x0"]))
    lines, cur = [], []
    for b in items:
        cy = (b["y0"] + b["y1"]) / 2
        if cur:
            ref = cur[0]
            rcy = (ref["y0"] + ref["y1"]) / 2
            if abs(cy - rcy) > 0.55 * max(ref["y1"] - ref["y0"], b["y1"] - b["y0"]):
                lines.append(cur)
                cur = []
        cur.append(b)
    if cur:
        lines.append(cur)
    out = []
    for ln in lines:
        ln.sort(key=lambda b: b["x0"])
        text = ln[0]["text"]
        for prev, b in zip(ln, ln[1:]):
            gap = b["x0"] - prev["x1"]
            hh = max(prev["y1"] - prev["y0"], 1)
            text += ("    " if gap > 2.5 * hh else " ") + b["text"]
        out.append({"x0": min(b["x0"] for b in ln), "x1": max(b["x1"] for b in ln),
                    "y0": min(b["y0"] for b in ln), "y1": max(b["y1"] for b in ln), "text": text,
                    "h": float(np.median([b["y1"] - b["y0"] for b in ln]))})
    return out


def _table_block(t, boxes):
    xs, ys = t["xs"], t["ys"]
    nr, nc = len(ys) - 1, len(xs) - 1
    grid = [[[] for _ in range(nc)] for _ in range(nr)]
    inside = [b for b in boxes if _center_in(b, t["bbox"])]
    for b in inside:
        cx, cy = (b["x0"] + b["x1"]) / 2, (b["y0"] + b["y1"]) / 2
        ci = max(0, min(nc - 1, int(np.searchsorted(xs, cx) - 1)))
        ri = max(0, min(nr - 1, int(np.searchsorted(ys, cy) - 1)))
        grid[ri][ci].append(b)
    cells = [["\n".join(l["text"] for l in group_lines(c)) for c in row] for row in grid]
    return {"type": "table", "y": t["bbox"][1], "cells": cells,
            "widths": [xs[i + 1] - xs[i] for i in range(nc)], "bbox": t["bbox"]}, inside


def build_blocks(arr, boxes):
    h, w = arr.shape[:2]
    # un vrai tableau contient du texte ; sinon (cadre, dessin) c'est une figure
    tables = [t for t in detect_tables(arr) if sum(_center_in(b, t["bbox"]) for b in boxes) >= 2]
    figs = detect_figures(arr, boxes, tables)
    blocks, used = [], set()
    for t in tables:
        blk, inside = _table_block(t, boxes)
        used.update(id(b) for b in inside)
        blocks.append(blk)
    for f in figs:
        used.update(id(b) for b in boxes if _center_in(b, f))
        blocks.append({"type": "image", "y": f[1], "bbox": f})
    text_boxes = [b for b in boxes if id(b) not in used]
    lines = group_lines(text_boxes)
    if lines:
        med = float(np.median([l["h"] for l in lines])) or 1.0
        left = min(l["x0"] for l in lines)
        right = max(l["x1"] for l in lines)
        mid = (left + right) / 2

        def align(l):
            cx = (l["x0"] + l["x1"]) / 2
            if (abs(cx - mid) < 0.04 * w or abs(cx - w / 2) < 0.05 * w) and (l["x0"] - left) > 0.08 * w:
                return "center"
            if (right - l["x1"]) < 0.03 * w and (l["x0"] - left) > 0.3 * w:
                return "right"
            return "left"

        para = None
        for l in lines:
            size = max(8.0, min(30.0, round(11 * l["h"] / med)))
            al = align(l)
            if (para and para["align"] == al and abs(para["size"] - size) <= 1 and
                    0 <= l["y0"] - para["y1"] < 0.9 * med and al == "left"
                    and abs(l["x0"] - para["x0"]) < 0.06 * w and not para["text"].endswith(("    ",))):
                para["text"] += " " + l["text"]
                para["y1"] = l["y1"]
                continue
            if para:
                blocks.append(para)
            para = {"type": "para", "y": l["y0"], "y1": l["y1"], "x0": l["x0"], "text": l["text"],
                    "size": size, "align": al,
                    "bold": size >= 14 or (l["text"].isupper() and len(l["text"]) < 70),
                    "gap": (l["y0"] - blocks[-1]["y1"] if blocks and "y1" in blocks[-1] else 0) / med}
        if para:
            blocks.append(para)
    blocks.sort(key=lambda b: b["y"])
    return blocks, len(tables), len(figs)


# ------------------------------------------------------------------- analyse
def _jpg(arr, q=88):
    import cv2
    ok, buf = cv2.imencode(".jpg", arr, [cv2.IMWRITE_JPEG_QUALITY, q])
    return buf.tobytes()


def _decode(jpg):
    import cv2
    return cv2.imdecode(np.frombuffer(jpg, np.uint8), cv2.IMREAD_COLOR)


def _load_image(path):
    from PIL import Image, ImageOps
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im).convert("RGB")
        if max(im.size) > MAX_SIDE:
            im.thumbnail((MAX_SIDE, MAX_SIDE))
        return np.ascontiguousarray(np.array(im)[:, :, ::-1])


def _analyse_array(arr, dpi, res):
    boxes = ocr_boxes(arr)
    blocks, nt, nf = build_blocks(arr, boxes)
    res.n_tables += nt
    h, w = arr.shape[:2]
    res.pages.append(Page(_jpg(arr), w, h, dpi, boxes, blocks))
    import cv2
    n = 0
    for b in blocks:
        if b["type"] == "image":
            x0, y0, x1, y1 = [int(v) for v in b["bbox"]]
            ok, buf = cv2.imencode(".png", arr[y0:y1, x0:x1])
            if ok:
                n += 1
                res.figures.append({"name": f"page{len(res.pages)}_image{n}.png", "data": buf.tobytes(), "page": len(res.pages)})
    res.n_figures = len(res.figures)


def _plain_text(page):
    parts = []
    for b in page.blocks:
        if b["type"] == "para":
            parts.append(b["text"])
        elif b["type"] == "table":
            parts.append("\n".join(" | ".join(c.replace("\n", " ") for c in row) for row in b["cells"]))
    return "\n\n".join(parts)


def clean_text(text):
    text = text.replace("\r", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(l.strip() for l in text.split("\n")).strip()


def _pdf_embedded_images(pdf, total):
    """Images réellement incorporées dans un PDF numérique (hors fonds de page entiers et miniatures)."""
    out, seen = [], set()
    for i in range(total):
        page = pdf[i]
        pr = page.rect
        for info in page.get_image_info(xrefs=True):
            xref = info.get("xref")
            if not xref or xref in seen:
                continue
            bb = info["bbox"]
            if info.get("width", 0) < 48 or info.get("height", 0) < 48:
                continue
            if (bb[2] - bb[0]) * (bb[3] - bb[1]) > 0.9 * pr.width * pr.height:
                continue
            try:
                img = pdf.extract_image(xref)
            except Exception:
                continue
            if not img or not img.get("image"):
                continue
            seen.add(xref)
            out.append({"name": f"page{i + 1}_image{len(out) + 1}.{img.get('ext', 'png')}", "data": img["image"], "page": i + 1})
    return out


def analyze(path, progress=None) -> OcrResult:
    p = Path(path)
    if not p.is_file():
        raise ValueError("Fichier introuvable.")
    if p.stat().st_size > MAX_BYTES:
        raise ValueError("Fichier trop volumineux (60 Mo maximum).")
    ext = p.suffix.lower()
    if ext in IMAGE_EXT:
        res = OcrResult(str(p), "image", n_pages=1)
        _analyse_array(_load_image(p), 150, res)
        res.text = clean_text(_plain_text(res.pages[0]))
        return res
    if ext != ".pdf":
        raise ValueError("Format non pris en charge (PNG, JPG, TIFF, PDF...).")
    with fitz.open(p) as pdf:
        if pdf.needs_pass:
            raise ValueError("Ce PDF est protégé par mot de passe : retirez-le d'abord (onglet PDF → Sécuriser).")
        total = min(len(pdf), MAX_PAGES)
        texts = [pdf[i].get_text().strip() for i in range(total)]
        res = OcrResult(str(p), "pdf_text", n_pages=total)
        if all(len(t) >= 30 for t in texts):
            res.text = clean_text("\n\n".join(texts))
            res.figures = _pdf_embedded_images(pdf, total)
            res.n_figures = len(res.figures)
            res.n_tables = sum(len(pdf[i].find_tables().tables) for i in range(total)) if hasattr(
                pdf[0], "find_tables") else 0
            return res
        res.kind = "pdf_scan"
        for i in range(total):
            pix = pdf[i].get_pixmap(matrix=fitz.Matrix(PDF_ZOOM, PDF_ZOOM), alpha=False)
            arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3)
            _analyse_array(np.ascontiguousarray(arr[:, :, ::-1]), 72 * PDF_ZOOM, res)
            if progress:
                progress(i + 1, total)
    res.text = clean_text("\n\n".join(_plain_text(pg) for pg in res.pages))
    return res


# -------------------------------------------------------------------- exports
def export_pdf(res: OcrResult, out):
    """PDF identique à l'original, avec une couche de texte invisible (recherche / copier)."""
    out = Path(out)
    if res.kind == "pdf_text":
        shutil.copyfile(res.path, out)
        return out
    doc = fitz.open()
    for pg in res.pages:
        s = 72.0 / pg.dpi
        page = doc.new_page(width=pg.w * s, height=pg.h * s)
        page.insert_image(page.rect, stream=pg.jpg)
        for b in pg.boxes:
            hh = (b["y1"] - b["y0"]) * s
            try:
                page.insert_text((b["x0"] * s, b["y1"] * s - 0.2 * hh), b["text"],
                                 fontsize=max(4, hh * 0.85), fontname="helv", render_mode=3)
            except Exception:
                continue
    doc.save(str(out), garbage=3, deflate=True)
    doc.close()
    return out


def export_docx(res: OcrResult, out):
    out = Path(out)
    if res.kind == "pdf_text":
        try:
            from pdf2docx import Converter
            cv = Converter(res.path)
            try:
                cv.convert(str(out))
            finally:
                cv.close()
            return out
        except Exception:
            from docx import Document
            d = Document()
            for para in res.text.split("\n\n"):
                d.add_paragraph(para)
            d.save(str(out))
            return out
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Pt
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)
    usable = 17.0
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    AL = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
    for pi, pg in enumerate(res.pages):
        if pi:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        arr = None
        for b in pg.blocks:
            if b["type"] == "para":
                par = doc.add_paragraph()
                par.alignment = AL[b["align"]]
                par.paragraph_format.space_after = Pt(4 if b.get("gap", 0) < 1.2 else 10)
                run = par.add_run(b["text"])
                run.font.size = Pt(b["size"])
                run.bold = b["bold"]
            elif b["type"] == "table":
                rows, cols = len(b["cells"]), len(b["cells"][0])
                tb = doc.add_table(rows=rows, cols=cols)
                tb.style = "Table Grid"
                tb.autofit = False
                tot = sum(b["widths"]) or 1
                for ri, row in enumerate(b["cells"]):
                    for ci, txt in enumerate(row):
                        cell = tb.cell(ri, ci)
                        cell.width = Cm(max(1.2, usable * b["widths"][ci] / tot))
                        cell.text = ""
                        run = cell.paragraphs[0].add_run(txt)
                        run.font.size = Pt(10)
                doc.add_paragraph()
            elif b["type"] == "image":
                if arr is None:
                    arr = _decode(pg.jpg)
                import cv2
                x0, y0, x1, y1 = [int(v) for v in b["bbox"]]
                ok, buf = cv2.imencode(".png", arr[y0:y1, x0:x1])
                width = max(1.5, min(usable, usable * (x1 - x0) / pg.w))
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER if abs(
                    (x0 + x1) / 2 - pg.w / 2) < 0.1 * pg.w else WD_ALIGN_PARAGRAPH.LEFT
                par.add_run().add_picture(io.BytesIO(buf.tobytes()), width=Cm(width))
    doc.save(str(out))
    return out
