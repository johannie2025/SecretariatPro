"""Rédaction assistée : documents types, mise en forme automatique du texte, export Word/PDF."""
import io
import re
from datetime import date
from decimal import Decimal
from pathlib import Path

from .billing import money
from .numwords import amount_in_words

MONTHS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre",
          "octobre", "novembre", "décembre"]
NB = "\u00a0"
STYLES = {"Classique (Times 12)": ("Times New Roman", 12, 1.15), "Moderne (Calibri 11)": ("Calibri", 11, 1.15),
          "Compact (Arial 10)": ("Arial", 10, 1.0)}


MONTHS_ES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
             "noviembre", "diciembre"]
MONTHS_EN = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
             "November", "December"]


def fr_date(d: date | None = None, lang: str = "fr") -> str:
    d = d or date.today()
    if lang == "es":
        return f"{d.day} de {MONTHS_ES[d.month - 1]} de {d.year}"
    if lang == "en":
        return f"{MONTHS_EN[d.month - 1]} {d.day}, {d.year}"
    return f"{d.day}{'er' if d.day == 1 else ''} {MONTHS[d.month - 1]} {d.year}"


# =============================================================== typographie
_ABBR = {"etc", "m", "mme", "mlle", "dr", "pr", "me", "st", "n", "tel", "cf", "ex", "vs", "art", "no"}


def typographie(text: str, lang: str = "fr") -> str:
    if lang != "fr":
        return _typo_latin(text, lang)
    t = text.replace("\u00a0", " ").replace("\u202f", " ")
    t = re.sub(r"[ \t]+", " ", t)
    t = t.replace("...", "…")
    t = re.sub(r"(?<=\w)'(?=\w)", "’", t)
    t = re.sub(r'"\s*([^"\n]+?)\s*"', f"«{NB}\\1{NB}»", t)
    t = re.sub(r"\s+([,.)\]…])", r"\1", t)                 # pas d'espace avant , . ) …
    t = re.sub(r"(\()\s+", r"\1", t)
    t = re.sub(r"\s*([;:!?])", lambda m: NB + m.group(1), t)   # espace insécable avant ; : ! ?
    t = re.sub(r"(?<=\d)" + NB + r":" + NB + r"?(?=\d)", ":", t)     # heures 10:30
    t = re.sub(r"(https?)" + NB + ":", r"\1:", t)                   # liens
    t = re.sub(r"«\s*", f"«{NB}", t)
    t = re.sub(r"\s*»", f"{NB}»", t)
    t = re.sub(r"(?<=[a-zà-ÿ])([,;:!?])(?=[A-Za-zÀ-ÿ])", r"\1 ", t)
    t = re.sub(r"(?<=[a-zà-ÿ]{2})\.(?=[A-ZÀ-Ý][a-zà-ÿ])", ". ", t)
    return t


def _typo_latin(text: str, lang: str) -> str:
    """Espagnol / anglais : pas d'espace avant la ponctuation, guillemets typographiques, apostrophes courbes."""
    t = text.replace("\u00a0", " ").replace("\u202f", " ")
    t = re.sub(r"[ \t]+", " ", t).replace("...", "…")
    t = re.sub(r"(?<=\w)'(?=\w)", "’", t)
    q1, q2 = ("«", "»") if lang == "es" else ("“", "”")
    t = re.sub(r'"\s*([^"\n]+?)\s*"', f"{q1}\\1{q2}", t)
    t = re.sub(r"\s+([,.;:!?)\]…»”])", r"\1", t)
    t = re.sub(r"([(\[«“¿¡])\s+", r"\1", t)
    t = re.sub(r"(?<=[a-zà-ÿ])([,;:!?])(?=[A-Za-zÀ-ÿ¿¡])", r"\1 ", t)
    t = re.sub(r"(?<=[a-zà-ÿ]{2})\.(?=[A-ZÀ-Ý][a-zà-ÿ])", ". ", t)
    return t


def _cap_sentences(text: str) -> str:
    def fix(m):
        before, letter = m.group(1), m.group(2)
        word = re.findall(r"([A-Za-zÀ-ÿ]+)[.!?…]\s*$", before)
        if word and word[-1].lower() in _ABBR:
            return m.group(0)
        return before + letter.upper()
    text = re.sub(r"((?:^|[.!?]\s+|\n)[«“¿¡\u00a0\s]*)([a-zà-ÿ])", fix, text, flags=re.M)
    return text


# ============================================================ mise en forme auto
_BULLET = re.compile(r"^\s*(?:[-–—•*·▪●]|\(?[a-z]\)|\d{1,2}[.)])\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d{1,2}[.)]\s+")
_HEADING_WORDS = re.compile(r"^(article|chapitre|titre|section|partie|annexe|objet)\b", re.I)


def format_text(raw: str, with_typo=True, lang: str = "fr") -> list:
    """Texte brut / issu d'un OCR → éléments structurés (titre, sous-titres, paragraphes, listes)."""
    raw = raw.replace("\r", "")
    raw = re.sub(r"(\w)-\n(\w)", r"\1\2", raw)
    tables = []

    def _grab_table(m):                       # tableaux Markdown « | a | b | » (issus de l'OCR IA)
        rows = []
        for ln in m.group(0).strip().split("\n"):
            if re.fullmatch(r"[\s|:\-]+", ln) and "-" in ln:
                continue
            rows.append([c.strip() for c in ln.strip().strip("|").split("|")])
        tables.append(rows)
        return f"\n\n\x00TABLE{len(tables) - 1}\x00\n\n"

    raw = re.sub(r"(?:^[ \t]*\|.*\|[ \t]*$\n?)+", _grab_table, raw, flags=re.M)
    blocks = [b for b in re.split(r"\n\s*\n", raw) if b.strip()]
    els, first = [], True

    def add_text(text):
        nonlocal first
        for para in text.split("\n"):
            para = para.strip()
            if not para:
                continue
            if para.startswith("\x00TABLE"):
                els.append({"t": "table", "rows": tables[int(re.findall(r"\d+", para)[0])]})
                first = False
                continue
            short = len(para) < 80 and not re.search(r"[.!?;,]$", para)
            if first and short and (para.isupper() or len(blocks) > 1):
                els.append({"t": "title", "text": para})
            elif short and (para.isupper() or _HEADING_WORDS.match(para) or re.match(r"^([IVX]+|\d+)[.)]\s+\S", para)):
                els.append({"t": "h", "text": para})
            else:
                els.append({"t": "p", "text": para})
            first = False

    for blk in blocks:
        text, items = "", []
        for l in [x.rstrip() for x in blk.split("\n") if x.strip()]:
            if _BULLET.match(l):
                if text:
                    add_text(text)
                    text = ""
                items.append(_BULLET.match(l).group(1).strip())
                numbered = _NUMBERED.match(l) is not None
                continue
            if items:
                els.append({"t": "olist" if numbered else "bullets", "items": items})
                items = []
            l = l.strip()
            if text and not re.search(r"[.!?…:;»]$", text):   # recoller les lignes coupées (OCR, copier-coller)
                text += " " + l
            else:
                text += ("\n" if text else "") + l
        if items:
            els.append({"t": "olist" if numbered else "bullets", "items": items})
        if text:
            add_text(text)
        first = False
    if with_typo:
        for e in els:
            if "text" in e:
                e["text"] = _cap_sentences(typographie(e["text"], lang))
            if "items" in e:
                e["items"] = [_cap_sentences(typographie(i, lang)) for i in e["items"]]
    return els


def elements_to_text(els) -> str:
    out = []
    for e in els:
        if e["t"] in ("title", "h", "p", "right", "center", "sig"):
            out.append(e["text"])
        elif e["t"] == "bullets":
            out.append("\n".join("• " + i for i in e["items"]))
        elif e["t"] == "olist":
            out.append("\n".join(f"{n}. {i}" for n, i in enumerate(e["items"], 1)))
        elif e["t"] == "letterhead":
            out.append("\n".join([e["name"]] + e["lines"]))
        elif e["t"] == "table":
            out.append("\n".join(" | ".join(r) for r in e["rows"]))
    return "\n\n".join(out)


# =============================================================== documents types
def _letterhead(sender):
    if not sender:
        return []
    lines = [x for x in (sender.get("address"), " — ".join(
        y for y in (sender.get("phone"), sender.get("email")) if y),
        " — ".join(y for y in (f"NIF : {sender['nif']}" if sender.get("nif") else "",
                                f"RCCM : {sender['rccm']}" if sender.get("rccm") else "") if y)) if x]
    return [{"t": "letterhead", "name": sender["name"], "lines": lines, "logo": sender.get("logo_path", "")}]


# Phrases des documents types, par langue (les accolades sont remplies par les builders).
PH = {
    "fr": {
        "subject": "Objet : {x}", "dear": "Madame, Monsieur,", "regards": "Veuillez agréer, Madame, Monsieur, l'expression de mes salutations distinguées.",
        "sign": "Signature", "placedate": "{lieu}, le {date}", "dear_conv": "Madame, Monsieur",
        "valid": "La présente attestation est délivrée pour servir et valoir ce que de droit.",
        "work_title": "ATTESTATION DE TRAVAIL", "work_sign": "Le responsable",
        "work": "Je soussigné(e), <b>{sig}</b>, {fn}, atteste que <b>{nom}</b>, né(e) le {nais} à {lnais}, est employé(e) en qualité de <b>{poste}</b> depuis le {debut}",
        "work_end": " jusqu'au {fin}.", "work_open": " et occupe toujours ce poste.", "cert": "ATTESTATION",
        "conv_title": "CONVOCATION", "conv": "Vous êtes convoqué(e) à <b>{obj}</b> le <b>{jour}</b> à <b>{heure}</b>, à l'adresse suivante : {adr}.",
        "agenda": "Ordre du jour :", "attend": "Votre présence est vivement souhaitée.",
        "pv_title": "PROCÈS-VERBAL DE RÉUNION", "pv_head": "<b>Objet :</b> {obj}<br/><b>Date :</b> {jour} à {heure}<br/><b>Lieu :</b> {lieu}",
        "attendees": "Participants", "proceed": "Déroulement et décisions",
        "pv_end": "La séance est levée à l'heure indiquée ci-dessus. Le présent procès-verbal est établi pour servir ce que de droit.",
        "pv_sign": "Le secrétaire de séance",
        "memo_title": "NOTE DE SERVICE N° {n}", "memo_head": "<b>De :</b> {de}<br/><b>À :</b> {a}<br/><b>Objet :</b> {obj}",
        "all_staff": "Tout le personnel", "memo_sign": "La direction",
        "rcpt_title": "REÇU DE PAIEMENT", "rcpt_no": " N° {n}", "amount": "Montant : {x}",
        "rcpt": "Reçu de <b>{de}</b> la somme de <b>{words}</b> ({shown} {cur}){motif}.", "motif": " pour : {m}",
        "pay_mode": "Mode de paiement : {x}.", "cash": "Espèces", "rcpt_sign": "Le bénéficiaire",
        "dis_title": "DÉCHARGE", "dis": "Je soussigné(e) <b>{sig}</b> reconnais avoir reçu de <b>{de}</b> : {obj}.",
        "dis_end": "La présente décharge est établie pour servir et valoir ce que de droit.",
        "dis_sign": "Signature précédée de « Lu et approuvé »",
        "poa_title": "PROCURATION", "poa": "Je soussigné(e) <b>{sig}</b>, demeurant à {adr}, donne par la présente pouvoir à <b>{mand}</b> afin de {obj}.",
        "poa2": "En conséquence, le mandataire pourra signer tous documents, faire toutes déclarations et, d'une manière générale, faire le nécessaire pour l'exécution du présent mandat.",
        "res_subject": "Démission",
        "res": ("Par la présente, j'ai l'honneur de vous informer de ma décision de démissionner de mon poste de {poste} au sein de votre structure.\n"
                "Conformément à mes obligations, j'effectuerai mon préavis de {preavis}, la fin de mon contrat étant fixée au {fin}.\n"
                "Je vous remercie pour la confiance accordée pendant toute cette période."),
        "rem_subject": "Relance — facture n° {f}", "delay": "8 jours",
        "rem": ("Sauf erreur de notre part, la facture n° {f} d'un montant de {m} {cur}, échue le {ech}, demeure impayée à ce jour.\n"
                "Nous vous prions de bien vouloir procéder à son règlement dans un délai de {delai}. À défaut, nous nous verrons contraints d'engager les démarches nécessaires au recouvrement.\n"
                "Si ce règlement a déjà été effectué, nous vous remercions de ne pas tenir compte de ce courrier."),
    },
    "es": {
        "subject": "Asunto: {x}", "dear": "Estimados señores:", "regards": "Sin otro particular, les saluda atentamente.",
        "sign": "Firma", "placedate": "{lieu}, a {date}", "dear_conv": "Estimados señores",
        "valid": "Se expide el presente certificado para los fines que estime oportunos.",
        "work_title": "CERTIFICADO DE TRABAJO", "work_sign": "El responsable",
        "work": "Quien suscribe, <b>{sig}</b>, {fn}, certifica que <b>{nom}</b>, nacido/a el {nais} en {lnais}, trabaja como <b>{poste}</b> desde el {debut}",
        "work_end": " hasta el {fin}.", "work_open": " y continúa ocupando este puesto.", "cert": "CERTIFICADO",
        "conv_title": "CONVOCATORIA", "conv": "Se le convoca a <b>{obj}</b> el <b>{jour}</b> a las <b>{heure}</b>, en la siguiente dirección: {adr}.",
        "agenda": "Orden del día:", "attend": "Se ruega encarecidamente su asistencia.",
        "pv_title": "ACTA DE REUNIÓN", "pv_head": "<b>Asunto:</b> {obj}<br/><b>Fecha:</b> {jour} a las {heure}<br/><b>Lugar:</b> {lieu}",
        "attendees": "Participantes", "proceed": "Desarrollo y decisiones",
        "pv_end": "Se levanta la sesión a la hora indicada. La presente acta se extiende para los fines oportunos.",
        "pv_sign": "El secretario de actas",
        "memo_title": "NOTA INTERNA N.º {n}", "memo_head": "<b>De:</b> {de}<br/><b>Para:</b> {a}<br/><b>Asunto:</b> {obj}",
        "all_staff": "Todo el personal", "memo_sign": "La dirección",
        "rcpt_title": "RECIBO DE PAGO", "rcpt_no": " N.º {n}", "amount": "Importe: {x}",
        "rcpt": "Recibí de <b>{de}</b> la suma de <b>{words}</b> ({shown} {cur}){motif}.", "motif": " por concepto de: {m}",
        "pay_mode": "Forma de pago: {x}.", "cash": "Efectivo", "rcpt_sign": "El beneficiario",
        "dis_title": "ACUSE DE RECIBO", "dis": "Quien suscribe, <b>{sig}</b>, reconoce haber recibido de <b>{de}</b>: {obj}.",
        "dis_end": "Se firma el presente documento para los fines oportunos.",
        "dis_sign": "Firma precedida de «Leído y aprobado»",
        "poa_title": "PODER", "poa": "Quien suscribe, <b>{sig}</b>, con domicilio en {adr}, otorga por la presente poder a <b>{mand}</b> para {obj}.",
        "poa2": "En consecuencia, el apoderado podrá firmar toda clase de documentos, efectuar declaraciones y, en general, hacer lo necesario para la ejecución de este mandato.",
        "res_subject": "Renuncia",
        "res": ("Por la presente, tengo el honor de comunicarle mi decisión de renunciar a mi puesto de {poste} en su organización.\n"
                "Cumpliré el preaviso de {preavis}; mi contrato finalizará el {fin}.\n"
                "Le agradezco la confianza depositada durante todo este tiempo."),
        "rem_subject": "Recordatorio de pago — factura n.º {f}", "delay": "8 días",
        "rem": ("Salvo error por nuestra parte, la factura n.º {f} por importe de {m} {cur}, vencida el {ech}, sigue pendiente de pago.\n"
                "Le rogamos que proceda a su abono en un plazo de {delai}. En caso contrario, nos veremos obligados a iniciar las gestiones de cobro necesarias.\n"
                "Si ya realizó este pago, le agradecemos que no tenga en cuenta esta carta."),
    },
    "en": {
        "subject": "Subject: {x}", "dear": "Dear Sir or Madam,", "regards": "Yours faithfully,",
        "sign": "Signature", "placedate": "{lieu}, {date}", "dear_conv": "Dear Sir or Madam",
        "valid": "This certificate is issued for whatever purpose it may serve.",
        "work_title": "EMPLOYMENT CERTIFICATE", "work_sign": "The manager",
        "work": "I, the undersigned, <b>{sig}</b>, {fn}, hereby certify that <b>{nom}</b>, born on {nais} in {lnais}, has been employed as <b>{poste}</b> since {debut}",
        "work_end": " until {fin}.", "work_open": " and still holds this position.", "cert": "CERTIFICATE",
        "conv_title": "NOTICE OF MEETING", "conv": "You are invited to <b>{obj}</b> on <b>{jour}</b> at <b>{heure}</b>, at the following address: {adr}.",
        "agenda": "Agenda:", "attend": "Your attendance is kindly requested.",
        "pv_title": "MINUTES OF MEETING", "pv_head": "<b>Subject:</b> {obj}<br/><b>Date:</b> {jour} at {heure}<br/><b>Venue:</b> {lieu}",
        "attendees": "Attendees", "proceed": "Proceedings and decisions",
        "pv_end": "The meeting was adjourned at the time stated above. These minutes are drawn up for the record.",
        "pv_sign": "The secretary",
        "memo_title": "INTERNAL MEMO No. {n}", "memo_head": "<b>From:</b> {de}<br/><b>To:</b> {a}<br/><b>Subject:</b> {obj}",
        "all_staff": "All staff", "memo_sign": "Management",
        "rcpt_title": "PAYMENT RECEIPT", "rcpt_no": " No. {n}", "amount": "Amount: {x}",
        "rcpt": "Received from <b>{de}</b> the sum of <b>{words}</b> ({shown} {cur}){motif}.", "motif": " for: {m}",
        "pay_mode": "Payment method: {x}.", "cash": "Cash", "rcpt_sign": "The recipient",
        "dis_title": "ACKNOWLEDGEMENT OF RECEIPT", "dis": "I, the undersigned, <b>{sig}</b>, acknowledge having received from <b>{de}</b>: {obj}.",
        "dis_end": "This document is issued for whatever purpose it may serve.",
        "dis_sign": "Signature preceded by \"Read and approved\"",
        "poa_title": "POWER OF ATTORNEY", "poa": "I, the undersigned, <b>{sig}</b>, residing at {adr}, hereby grant power of attorney to <b>{mand}</b> in order to {obj}.",
        "poa2": "Accordingly, the attorney may sign any documents, make any declarations and, in general, do whatever is necessary to carry out this mandate.",
        "res_subject": "Resignation",
        "res": ("I am writing to inform you of my decision to resign from my position as {poste} with your organisation.\n"
                "In accordance with my obligations, I will work my notice period of {preavis}, my contract ending on {fin}.\n"
                "Thank you for the trust placed in me throughout this time."),
        "rem_subject": "Payment reminder — invoice no. {f}", "delay": "8 days",
        "rem": ("Unless we are mistaken, invoice no. {f} for {m} {cur}, due on {ech}, remains unpaid.\n"
                "We kindly ask you to settle it within {delai}. Otherwise we will have to take the necessary steps to recover the amount.\n"
                "If you have already made this payment, please disregard this letter."),
    },
}


def _L(v):
    return PH.get(v.get("_lang", "fr"), PH["fr"])


def _place_date(v):
    return {"t": "right", "text": _L(v)["placedate"].format(lieu=v.get("lieu", "").strip() or "…", date=v.get("date") or fr_date(None, v.get("_lang", "fr")))}


def _paras(text, lang="fr"):
    return [{"t": "p", "text": typographie(p.strip(), lang)} for p in re.split(r"\n+", text or "") if p.strip()]


def _sig(v, label=None):
    return [{"t": "space", "n": 1}, {"t": "right", "text": label or _L(v)["sign"]}, {"t": "space", "n": 2},
            {"t": "right", "text": v.get("signataire", "")}]


def b_lettre(v, s):
    L = _L(v)
    e = _letterhead(s)
    if not s and v.get("expediteur"):
        e += [{"t": "left", "text": v["expediteur"]}]
    e += [_place_date(v),
          {"t": "right", "text": "\n".join(x for x in (v.get("destinataire"), v.get("fonction"), v.get("adresse")) if x)},
          {"t": "p", "text": "<b>" + L["subject"].format(x=v.get("objet", "")) + "</b>"}, {"t": "p", "text": L["dear"]}]
    e += _paras(v.get("corps"), v["_lang"]) + [{"t": "p", "text": v.get("formule") or L["regards"]}]
    return e + _sig(v)


def b_attestation_travail(v, s):
    L = _L(v)
    end = L["work_end"].format(fin=v["fin"]) if v.get("fin") else L["work_open"]
    return _letterhead(s) + [
        {"t": "title", "text": L["work_title"]},
        {"t": "p", "text": L["work"].format(sig=v.get("signataire", ""), fn=v.get("fonction_sign", ""), nom=v.get("nom", ""),
                                            nais=v.get("naissance", "…"), lnais=v.get("lieu_naissance", "…"),
                                            poste=v.get("poste", ""), debut=v.get("debut", "…")) + end},
        {"t": "p", "text": L["valid"]}, _place_date(v)] + _sig(v, L["work_sign"])


def b_attestation(v, s):
    L = _L(v)
    return _letterhead(s) + [{"t": "title", "text": (v.get("titre") or L["cert"]).upper()}] + _paras(v.get("corps"), v["_lang"]) + [
        {"t": "p", "text": L["valid"]}, _place_date(v)] + _sig(v)


def b_convocation(v, s):
    L = _L(v)
    return _letterhead(s) + [_place_date(v), {"t": "title", "text": L["conv_title"]},
        {"t": "p", "text": "<b>" + L["subject"].format(x=v.get("objet", "")) + "</b>"},
        {"t": "p", "text": f"{v.get('destinataires') or L['dear_conv']},"},
        {"t": "p", "text": L["conv"].format(obj=v.get("objet", "…"), jour=v.get("jour", "…"), heure=v.get("heure", "…"), adr=v.get("adresse", "…"))},
        {"t": "p", "text": "<b>" + L["agenda"] + "</b>"},
        {"t": "bullets", "items": [x.strip() for x in (v.get("ordre") or "").split("\n") if x.strip()] or ["…"]},
        {"t": "p", "text": L["attend"]}] + _sig(v)


def b_pv(v, s):
    L = _L(v)
    return _letterhead(s) + [{"t": "title", "text": L["pv_title"]},
        {"t": "p", "text": L["pv_head"].format(obj=v.get("objet", ""), jour=v.get("jour", ""), heure=v.get("heure", ""), lieu=v.get("lieu", ""))},
        {"t": "h", "text": L["attendees"]},
        {"t": "bullets", "items": [x.strip() for x in (v.get("participants") or "").split("\n") if x.strip()] or ["…"]},
        {"t": "h", "text": L["proceed"]}] + _paras(v.get("contenu"), v["_lang"]) + [
        {"t": "p", "text": L["pv_end"]}] + _sig(v, L["pv_sign"])


def b_note(v, s):
    L = _L(v)
    return _letterhead(s) + [{"t": "title", "text": L["memo_title"].format(n=v.get("numero", "…"))},
        {"t": "p", "text": L["memo_head"].format(de=v.get("de", ""), a=v.get("a") or L["all_staff"], obj=v.get("objet", ""))},
        _place_date(v)] + _paras(v.get("corps"), v["_lang"]) + _sig(v, L["memo_sign"])


def b_recu(v, s):
    L = _L(v)
    lang = v["_lang"]
    cur = v.get("devise") or "FCFA"
    try:
        amt = Decimal(str(v.get("montant", "0")).replace(" ", "").replace(",", "."))
    except Exception:
        raise ValueError("Montant invalide.")
    shown = money(amt, cur, lang)[: -len(cur) - 1]
    motif = L["motif"].format(m=v["motif"]) if v.get("motif") else ""
    return _letterhead(s) + [{"t": "title", "text": L["rcpt_title"] + (L["rcpt_no"].format(n=v["numero"]) if v.get("numero") else "")},
        {"t": "p", "text": "<b>" + L["amount"].format(x=f"{shown} {cur}") + "</b>"},
        {"t": "p", "text": L["rcpt"].format(de=v.get("recu_de", ""), words=amount_in_words(amt, cur, lang), shown=shown, cur=cur, motif=motif)},
        {"t": "p", "text": L["pay_mode"].format(x=v.get("mode") or L["cash"])}, _place_date(v)] + _sig(v, L["rcpt_sign"])


def b_decharge(v, s):
    L = _L(v)
    return _letterhead(s) + [{"t": "title", "text": L["dis_title"]},
        {"t": "p", "text": L["dis"].format(sig=v.get("signataire", ""), de=v.get("de", ""), obj=v.get("objet", ""))},
        {"t": "p", "text": L["dis_end"]}, _place_date(v)] + _sig(v, L["dis_sign"])


def b_procuration(v, s):
    L = _L(v)
    return _letterhead(s) + [{"t": "title", "text": L["poa_title"]},
        {"t": "p", "text": L["poa"].format(sig=v.get("signataire", ""), adr=v.get("adresse", "…"), mand=v.get("mandataire", ""), obj=v.get("objet", ""))},
        {"t": "p", "text": L["poa2"]}, _place_date(v)] + _sig(v)


def b_demission(v, s):
    L = _L(v)
    v = dict(v, objet=L["res_subject"])
    v["corps"] = L["res"].format(poste=v.get("poste", "…"), preavis=v.get("preavis", "…"), fin=v.get("fin", "…"))
    return b_lettre(v, s)


def b_relance(v, s):
    L = _L(v)
    v = dict(v, objet=L["rem_subject"].format(f=v.get("facture", "…")))
    v["corps"] = L["rem"].format(f=v.get("facture", "…"), m=v.get("montant", "…"), cur=v.get("devise") or "FCFA",
                                 ech=v.get("echeance", "…"), delai=v.get("delai") or L["delay"])
    return b_lettre(v, s)


FALLBACK_TEXT = {
    "fr": {"corps": "Nous vous écrivons au sujet de « {objet} ».\n\nNous restons à votre disposition pour toute information complémentaire.",
           "ordre": "Ouverture de la séance\nPoints à traiter\nQuestions diverses"},
    "es": {"corps": "Le escribimos en relación con «{objet}».\n\nQuedamos a su disposición para cualquier información adicional.",
           "ordre": "Apertura de la sesión\nPuntos a tratar\nRuegos y preguntas"},
    "en": {"corps": "We are writing to you regarding \"{objet}\".\n\nWe remain at your disposal for any further information.",
           "ordre": "Opening of the meeting\nItems to discuss\nAny other business"},
}

COMMON = [("lieu", "Lieu", "line"), ("date", "Date (laisser vide = aujourd'hui)", "line"), ("signataire", "Nom du signataire", "line")]
DOC_TYPES = {
    "Lettre administrative": (b_lettre, [("expediteur", "Expéditeur (si pas d'en-tête)", "text"), ("destinataire", "Destinataire", "line"),
        ("fonction", "Fonction du destinataire", "line"), ("adresse", "Adresse du destinataire", "line"),
        ("objet", "Objet", "line"), ("corps", "Corps de la lettre (un paragraphe par ligne)", "text"),
        ("formule", "Formule de politesse (facultatif)", "line")] + COMMON),
    "Lettre de démission": (b_demission, [("destinataire", "Destinataire", "line"), ("adresse", "Adresse", "line"),
        ("poste", "Poste occupé", "line"), ("preavis", "Durée du préavis", "line"), ("fin", "Date de fin de contrat", "line")] + COMMON),
    "Relance de paiement": (b_relance, [("destinataire", "Client débiteur", "line"), ("adresse", "Adresse", "line"),
        ("facture", "N° de facture", "line"), ("montant", "Montant", "line"), ("devise", "Devise", "line"),
        ("echeance", "Date d'échéance", "line"), ("delai", "Délai accordé", "line")] + COMMON),
    "Attestation de travail": (b_attestation_travail, [("nom", "Nom de l'employé(e)", "line"), ("naissance", "Date de naissance", "line"),
        ("lieu_naissance", "Lieu de naissance", "line"), ("poste", "Poste", "line"), ("debut", "Date d'embauche", "line"),
        ("fin", "Date de fin (vide si en poste)", "line"), ("fonction_sign", "Fonction du signataire", "line")] + COMMON),
    "Attestation (libre)": (b_attestation, [("titre", "Titre", "line"), ("corps", "Texte", "text")] + COMMON),
    "Convocation": (b_convocation, [("destinataires", "Destinataire(s)", "line"), ("objet", "Objet / réunion", "line"),
        ("jour", "Jour", "line"), ("heure", "Heure", "line"), ("adresse", "Lieu", "line"),
        ("ordre", "Ordre du jour (un point par ligne)", "text")] + COMMON),
    "Procès-verbal de réunion": (b_pv, [("objet", "Objet", "line"), ("jour", "Date", "line"), ("heure", "Heure", "line"),
        ("participants", "Participants (un par ligne)", "text"), ("contenu", "Déroulement et décisions", "text")] + COMMON),
    "Note de service": (b_note, [("numero", "N°", "line"), ("de", "De", "line"), ("a", "À", "line"), ("objet", "Objet", "line"),
        ("corps", "Texte", "text")] + COMMON),
    "Reçu de paiement": (b_recu, [("numero", "N° du reçu", "line"), ("recu_de", "Reçu de", "line"), ("montant", "Montant", "line"),
        ("devise", "Devise (FCFA, EUR...)", "line"), ("motif", "Motif", "line"), ("mode", "Mode de paiement", "line")] + COMMON),
    "Décharge": (b_decharge, [("de", "Reçu de", "line"), ("objet", "Objet remis", "text")] + COMMON),
    "Procuration": (b_procuration, [("adresse", "Adresse du mandant", "line"), ("mandataire", "Mandataire", "line"),
        ("objet", "Objet de la procuration", "text")] + COMMON),
}


def build_document(doc_type, values, sender=None, lang="fr"):
    builder, _ = DOC_TYPES[doc_type]
    lang = lang if lang in PH else "fr"
    v = {k: (x.strip() if isinstance(x, str) else x) for k, x in values.items()}
    v["_lang"] = lang
    if not v.get("date"):
        v["date"] = fr_date(None, lang)
    return builder(v, sender)


# ================================================================ rendu Word / PDF
def _strip_tags(t):
    return re.sub(r"</?b>", "", t).replace("<br/>", "\n")


def _runs(par, text, size=None, bold=False):
    parts = re.split(r"(<b>.*?</b>)", text.replace("<br/>", "\n"))
    for part in parts:
        if not part:
            continue
        b = part.startswith("<b>")
        r = par.add_run(_strip_tags(part))
        r.bold = bold or b
        if size:
            r.font.size = size


def to_docx(els, path, style="Classique (Times 12)"):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH as A
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    font, size, spacing = STYLES.get(style, STYLES["Classique (Times 12)"])
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = s.bottom_margin = Cm(2.2)
    n = d.styles["Normal"]
    n.font.name, n.font.size = font, Pt(size)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    n.paragraph_format.line_spacing = spacing
    n.paragraph_format.space_after = Pt(6)
    for e in els:
        t = e["t"]
        if t == "letterhead":
            if e.get("logo") and Path(e["logo"]).is_file():
                try:
                    d.add_paragraph().add_run().add_picture(e["logo"], height=Cm(2.0))
                except Exception:
                    pass
            p = d.add_paragraph()
            p.alignment = A.CENTER
            _runs(p, e["name"], Pt(size + 4), True)
            for ln in e["lines"]:
                q = d.add_paragraph()
                q.alignment = A.CENTER
                q.paragraph_format.space_after = Pt(0)
                _runs(q, ln, Pt(size - 1))
            pPr = q._p.get_or_add_pPr() if e["lines"] else p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bt = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "6"), ("w:color", "1F5FBF")):
                bt.set(qn(k), v)
            bd.append(bt)
            pPr.append(bd)
            d.add_paragraph()
        elif t == "title":
            p = d.add_paragraph()
            p.alignment = A.CENTER
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(12), Pt(14)
            _runs(p, e["text"], Pt(size + 5), True)
        elif t == "h":
            p = d.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.keep_with_next = True
            _runs(p, e["text"], Pt(size + 1), True)
        elif t in ("p", "left", "right", "center"):
            p = d.add_paragraph()
            p.alignment = {"p": A.JUSTIFY, "left": A.LEFT, "right": A.RIGHT, "center": A.CENTER}[t]
            _runs(p, e["text"])
        elif t == "sig":
            p = d.add_paragraph()
            p.alignment = A.RIGHT
            _runs(p, e["text"])
        elif t == "bullets":
            for i in e["items"]:
                d.add_paragraph(i, style="List Bullet")
        elif t == "olist":
            for i in e["items"]:
                d.add_paragraph(i, style="List Number")
        elif t == "space":
            for _ in range(e.get("n", 1)):
                d.add_paragraph()
        elif t == "table":
            ncol = max(len(r) for r in e["rows"])
            tb = d.add_table(rows=len(e["rows"]), cols=ncol)
            tb.style = "Table Grid"
            for i, r in enumerate(e["rows"]):
                for j in range(ncol):
                    cell = tb.cell(i, j)
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(r[j] if j < len(r) else "")
                    run.font.size = Pt(max(8, size - 1))
                    run.bold = i == 0
            d.add_paragraph()
        elif t == "image":
            try:
                d.add_paragraph().add_run().add_picture(io.BytesIO(e["data"]), width=Cm(min(e.get("width", 12), 16)))
            except Exception:
                pass
    d.save(str(path))
    return Path(path)


def to_pdf(els, path, style="Classique (Times 12)"):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import HRFlowable, Image, Paragraph, SimpleDocTemplate, Spacer
    font, size, spacing = STYLES.get(style, STYLES["Classique (Times 12)"])
    base, bold = ("Times-Roman", "Times-Bold") if "Times" in font else ("Helvetica", "Helvetica-Bold")
    ps = lambda **k: ParagraphStyle("x", fontName=k.pop("fontName", base), fontSize=k.pop("fontSize", size),
                                    leading=k.pop("leading", size * spacing * 1.2), spaceAfter=k.pop("spaceAfter", 6), **k)
    esc = lambda t: t.replace("&", "&amp;").replace("<b>", "\x01").replace("</b>", "\x02").replace("<", "&lt;").replace(
        ">", "&gt;").replace("\x01", "<b>").replace("\x02", "</b>").replace("<br/>", "<br/>").replace("\n", "<br/>")
    esc = lambda t, _e=esc: _e(t.replace("<br/>", "\n")).replace("\u00a0", " ")
    story = []
    for e in els:
        t = e["t"]
        if t == "letterhead":
            if e.get("logo") and Path(e["logo"]).is_file():
                try:
                    from PIL import Image as PI
                    with PI.open(e["logo"]) as im:
                        w, h = im.size
                    story.append(Image(e["logo"], width=2 * cm * w / h, height=2 * cm, hAlign="CENTER"))
                except Exception:
                    pass
            story.append(Paragraph(f"<b>{esc(e['name'])}</b>", ps(alignment=TA_CENTER, fontSize=size + 4, leading=size + 8, spaceAfter=2)))
            for ln in e["lines"]:
                story.append(Paragraph(esc(ln), ps(alignment=TA_CENTER, fontSize=size - 1, spaceAfter=0)))
            story += [HRFlowable(width="100%", thickness=1, color=colors.HexColor("#1F5FBF"), spaceBefore=6, spaceAfter=14)]
        elif t == "title":
            story.append(Paragraph(f"<b>{esc(e['text'])}</b>", ps(alignment=TA_CENTER, fontSize=size + 5, leading=size + 10, spaceBefore=10, spaceAfter=14)))
        elif t == "h":
            story.append(Paragraph(f"<b>{esc(e['text'])}</b>", ps(fontSize=size + 1, spaceBefore=8)))
        elif t in ("p", "left", "right", "center", "sig"):
            al = {"p": TA_JUSTIFY, "left": TA_LEFT, "right": TA_RIGHT, "center": TA_CENTER, "sig": TA_RIGHT}[t]
            story.append(Paragraph(esc(e["text"]), ps(alignment=al)))
        elif t in ("bullets", "olist"):
            for k, i in enumerate(e["items"], 1):
                story.append(Paragraph(esc(i), ps(leftIndent=18, bulletIndent=4, spaceAfter=3),
                                       bulletText="•" if t == "bullets" else f"{k}."))
        elif t == "space":
            story.append(Spacer(1, e.get("n", 1) * 0.6 * cm))
        elif t == "table":
            from reportlab.platypus import Table, TableStyle
            cell = ps(fontSize=size - 1, leading=(size - 1) * 1.2, spaceAfter=0)
            ncol = max(len(r) for r in e["rows"])
            data = [[Paragraph(esc(r[j] if j < len(r) else ""), cell) for j in range(ncol)] for r in e["rows"]]
            tb = Table(data, repeatRows=1)
            tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF1FC"))]))
            story += [tb, Spacer(1, 8)]
        elif t == "image":
            try:
                from PIL import Image as PI
                with PI.open(io.BytesIO(e["data"])) as im:
                    w0, h0 = im.size
                wd = min(e.get("width", 12), 16) * cm
                story += [Image(io.BytesIO(e["data"]), width=wd, height=wd * h0 / w0), Spacer(1, 6)]
            except Exception:
                pass
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2.5 * cm, rightMargin=2.5 * cm, topMargin=2.2 * cm,
                      bottomMargin=2.2 * cm).build(story)
    return Path(path)
